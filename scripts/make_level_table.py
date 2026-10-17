"""Generate a level 1..N table of every calc expression for one class.

Usage: python scripts/make_level_table.py data/skills.csv pal out.xlsx [maxlvl]

Assumption: only the evaluated skill has hard points (= its level);
all synergy sources are 0. Values are engine outputs, not formulas.
"""
import math
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from d2r.calc import CalcSyntaxError, UnsupportedToken
from d2r.skill import EXPR_COLUMNS, TOKEN_STATUS, Character, SkillDB, SkillEvaluator


def label(row, col):
    def s(c):
        v = row.get(c)
        return v if isinstance(v, str) and v.strip() else None
    if col.startswith("calc"):
        n = col[4:]
        return s(f"*calc{n} desc") or s(f"*calc{n}desc")
    if col.startswith("aurastatcalc"):
        return s(f"aurastat{col[12:]}") or "(aurastat 이름 없음)"
    if col.startswith("passivecalc"):
        return s(f"passivestat{col[11:]}") or "(passivestat 이름 없음)"
    return {"auralencalc": "오라/버프 지속(frames)", "aurarangecalc": "오라 반경",
            "EDmgSymPerCalc": "원소 데미지 시너지 %", "DmgSymPerCalc": "물리 데미지 시너지 %",
            "ToHitCalc": "AR"}.get(col)


def main(src, cls, out, maxlvl=30):
    db = SkillDB.load(src)
    levels = list(range(1, maxlvl + 1))
    rows, errs = [], Counter()

    for sk in db.by_class(cls):
        r = db[sk]
        ev_at = lambda L: SkillEvaluator(db, Character(hard={sk: L}))  # noqa: E731
        for col in EXPR_COLUMNS:
            expr = r.get(col)
            if not isinstance(expr, str) or not expr.strip():
                continue
            vals = []
            for L in levels:
                try:
                    vals.append(ev_at(L).column(sk, col, L))
                except (UnsupportedToken, CalcSyntaxError) as e:
                    vals.append(f"ERR {e}")
                    errs[str(e)] += 1
            rows.append((sk, col, label(r, col), expr, vals))

        e = ev_at(1)
        if not math.isnan(float(r.get("MinDam") or "nan")):
            rows.append((sk, "MinDam..MaxLevDam5", "물리 데미지 (Level Damage Formula)", "level formula",
                         [f"{a}-{b}" for a, b in (ev_at(L).phys_damage(sk, L) for L in levels)]))
        if isinstance(r.get("EType"), str):
            rows.append((sk, "EMin..EMaxLev5", f"{r['EType']} 데미지 (HitShift {int(r['HitShift'])}, 시너지 0)",
                         "level formula", [f"{lo}-{hi}" for _, lo, hi in
                                           (ev_at(L).elem_damage(sk, L) for L in levels)]))
        if not math.isnan(float(r.get("ELen") or "nan")):
            rows.append((sk, "ELen..ELevLen3", "원소 지속(frames)", "level formula",
                         [e.elem_length(sk, L) for L in levels]))
        if not math.isnan(float(r.get("ToHit") or "nan")):
            rows.append((sk, "ToHit/LevToHit", "AR 보너스 %", "toht",
                         [ev_at(L).token(sk, "toht", L, L) for L in levels]))

    wb = Workbook()
    ws = wb.active
    ws.title = f"{cls} Lv1-{maxlvl}"
    hdr = ["스킬", "컬럼", "의미", "수식"] + [f"Lv{L}" for L in levels]
    ws.append(hdr)
    bold, fill = Font(name="Arial", bold=True), PatternFill("solid", fgColor="DDE3EA")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    base = Font(name="Arial")
    alt = PatternFill("solid", fgColor="F4F6F8")
    prev, shade = None, False
    for sk, col, lab, expr, vals in rows:
        if sk != prev:
            shade, prev = not shade, sk
        ws.append([sk, col, lab, expr] + vals)
        for c in ws[ws.max_row]:
            c.font = base
            if shade:
                c.fill = alt
            if isinstance(c.value, str) and c.value.startswith("ERR"):
                c.font = Font(name="Arial", color="C00000")
    ws.freeze_panes = "E2"
    for i, w in enumerate([20, 18, 30, 44] + [8] * len(levels), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws["A1"].comment = Comment(
        "엔진 출력값. 가정: 해당 스킬만 hard point = 레벨, 시너지 소스 0, +스킬 0.", "d2r_sim")

    t = wb.create_sheet("토큰 의미·확신도")
    t.append(["토큰", "상태 / 정의"])
    for k, v in TOKEN_STATUS.items():
        t.append([k, v])
    t.append([])
    t.append(["미해석 토큰", "건수"])
    for k, n in errs.most_common():
        t.append([k, n])
    for c in t[1]:
        c.font, c.fill = bold, fill
    t.column_dimensions["A"].width = 16
    t.column_dimensions["B"].width = 90
    wb.save(out)
    print(f"{len(rows)} rows, {sum(errs.values())} errors -> {out}")


if __name__ == "__main__":
    a = sys.argv[1:]
    main(a[0], a[1], a[2], int(a[3]) if len(a) > 3 else 30)
